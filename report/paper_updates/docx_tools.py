"""Small helpers for editing the paper / Blue Book copies by body-element index, keeping formatting."""
from __future__ import annotations

import copy
import struct

from docx import Document
from docx.oxml.ns import qn

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
EMU_PER_IN = 914400


def body(doc):
    return list(doc.element.body.iterchildren())


def text_of(el) -> str:
    return "".join(t.text or "" for t in el.iter(W + "t"))


def set_par_text(p_el, new_text: str, expect: str | None = None) -> None:
    """Replace a paragraph's text, keeping the first text run's formatting."""
    if expect is not None:
        assert expect in text_of(p_el), f"expected {expect!r} in {text_of(p_el)[:80]!r}"
    runs = [r for r in p_el.findall(W + "r") if r.find(W + "t") is not None]
    for hl in p_el.findall(W + "hyperlink"):
        runs += [r for r in hl.findall(W + "r") if r.find(W + "t") is not None]
    assert runs, "paragraph has no text runs"
    first = runs[0]
    for r in runs[1:]:
        r.getparent().remove(r)
    ts = first.findall(W + "t")
    for t in ts[1:]:
        first.remove(t)
    ts[0].text = new_text
    ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def insert_par_after(anchor_el, template_p_el, new_text: str):
    """Deep-copy `template_p_el` (for its formatting), set its text, insert after `anchor_el`."""
    new = copy.deepcopy(template_p_el)
    for d in list(new.iter(W + "drawing")):
        d.getparent().remove(d)
    set_par_text(new, new_text)
    anchor_el.addnext(new)
    return new


def cell_texts(tbl_el):
    return [["".join(t.text or "" for t in tc.iter(W + "t")) for tc in tr.findall(W + "tc")]
            for tr in tbl_el.findall(W + "tr")]


def set_cell(tbl_el, r: int, c: int, new_text: str, expect: str | None = None) -> None:
    tc = tbl_el.findall(W + "tr")[r].findall(W + "tc")[c]
    ps = tc.findall(W + "p")
    if expect is not None:
        assert expect in "".join(text_of(p) for p in ps), f"cell {r},{c}: {expect!r} not found"
    target = next((p for p in ps if p.find(f".//{W}t") is not None), ps[0])
    if target.find(f".//{W}t") is None:  # empty cell: add a run
        rr = target.makeelement(W + "r", {})
        t = rr.makeelement(W + "t", {})
        rr.append(t)
        target.append(rr)
    set_par_text(target, new_text)
    for p in ps:
        if p is not target:
            tc.remove(p)


def image_info(doc, p_el):
    blips = list(p_el.iter(A + "blip"))
    if not blips:
        return None
    rid = blips[0].get(R_EMBED)
    ext = next(p_el.iter(WP + "extent"))
    return rid, int(ext.get("cx")), int(ext.get("cy"))


def replace_image(doc, p_el, png_path: str, width_in: float | None = None) -> None:
    """Swap the picture in paragraph `p_el` for `png_path`, keeping its width (or `width_in`)."""
    rid, cx, cy = image_info(doc, p_el)
    if width_in is not None:
        cx = int(width_in * EMU_PER_IN)
        for ext in p_el.iter(WP + "extent"):
            ext.set("cx", str(cx))
        for ext in p_el.iter(A + "ext"):
            if ext.get("cx") is not None:
                ext.set("cx", str(cx))
    part = doc.part.related_parts[rid]
    data = open(png_path, "rb").read()
    part._blob = data
    w, h = struct.unpack(">II", data[16:24])  # PNG IHDR
    new_cy = int(cx * h / w)
    for ext in p_el.iter(WP + "extent"):
        ext.set("cy", str(new_cy))
    for ext in p_el.iter(A + "ext"):
        if ext.get("cx") is not None:
            ext.set("cy", str(new_cy))


def insert_image_after(doc, anchor_el, png_path: str, width_in: float):
    """New centred paragraph holding a picture, inserted after `anchor_el`."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(png_path, width=Inches(width_in))
    anchor_el.addnext(p._p)
    return p._p


def text_runs(p_el):
    runs = [r for r in p_el.iter(W + "r") if r.find(W + "t") is not None]
    return runs


def replace_in_runs(p_el, old: str, new: str) -> bool:
    """Replace `old` inside a single run, keeping every run's formatting. False if it spans runs."""
    for r in text_runs(p_el):
        ts = r.findall(W + "t")
        joined = "".join(t.text or "" for t in ts)
        if old in joined:
            for t in ts[1:]:
                r.remove(t)
            ts[0].text = joined.replace(old, new, 1)
            ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            return True
    return False


def insert_head_body_after(anchor_el, template_p_el, head: str, rest: str):
    """Copy a list item like '(5)  Living surveys. Because…' keeping bold head + normal body."""
    new = copy.deepcopy(template_p_el)
    runs = text_runs(new)
    first, last = runs[0], runs[-1]
    for r in runs[1:-1]:
        r.getparent().remove(r)
    for r, txt in ((first, head), (last, rest)):
        ts = r.findall(W + "t")
        for t in ts[1:]:
            r.remove(t)
        ts[0].text = txt
        ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    anchor_el.addnext(new)
    return new


def set_head_body(p_el, head: str, rest: str) -> None:
    """In place: first text run gets `head` (e.g. bold 'Fig. 4.'), last run gets `rest`; others removed."""
    runs = text_runs(p_el)
    if len(runs) < 2:
        set_par_text(p_el, head + rest)
        return
    first, last = runs[0], runs[-1]
    for r in runs[1:-1]:
        r.getparent().remove(r)
    for r, txt in ((first, head), (last, rest)):
        ts = r.findall(W + "t")
        for t in ts[1:]:
            r.remove(t)
        ts[0].text = txt
        ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
