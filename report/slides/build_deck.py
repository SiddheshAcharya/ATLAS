"""Build the final-review deck (report/slides/ATLAS_Final_Review.pptx) with python-pptx.

Same visual language as the team's Review I deck (16:9, Cambria navy titles, navy footer bar,
Calibri body, red/navy tag labels). Every number here comes from docs/TESTING.md /
report/report.md as of 2026-09-29 — edit the numbers there first, then here, then rebuild:
    .venv/Scripts/python report/slides/build_deck.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"
OUT = HERE / "ATLAS_Final_Review.pptx"

NAVY = RGBColor(0x0D, 0x12, 0x26)
GREY = RGBColor(0xAE, 0xB6, 0xCC)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x55, 0x60, 0x70)
RED = RGBColor(0xE2, 0x4B, 0x4A)
RED_BG = RGBColor(0xFC, 0xEB, 0xEB)
GREEN = RGBColor(0x1E, 0x8E, 0x5A)
GREEN_BG = RGBColor(0xE7, 0xF6, 0xEE)
BLUE = RGBColor(0x2F, 0x6F, 0xD6)
BLUE_BG = RGBColor(0xEA, 0xF1, 0xFD)
AMBER = RGBColor(0xB7, 0x79, 0x1F)
AMBER_BG = RGBColor(0xFD, 0xF4, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

AUTHORS = "Siddhesh Acharya   ·   Mrigakshi Kaul   ·   Suhani Singh"
FOOTER = ("ATLAS  —  Agentic Technology for Literature Analysis and Synthesis   ·   "
          "Siddhesh Acharya, Mrigakshi Kaul, Suhani Singh   ·   Final Review")

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=None, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    return shp


def text(slide, x, y, w, h, content, *, size=13, bold=False, color=INK, font="Calibri",
         align=PP_ALIGN.LEFT, italic=False):
    """`content`: a string (newlines = paragraphs) or a list of (text, overrides) runs-per-line."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = content if isinstance(content, list) else content.split("\n")
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = line if isinstance(line, list) else [(line, {})]
        for seg, over in segs:
            run = para.add_run()
            run.text = seg
            f = run.font
            f.name = over.get("font", font)
            f.size = Pt(over.get("size", size))
            f.bold = over.get("bold", bold)
            f.italic = over.get("italic", italic)
            f.color.rgb = over.get("color", color)
        para.space_after = Pt(4)
    return tb


def tag(slide, x, y, label, fill=RED, w=None):
    w = w or (0.13 + 0.085 * len(label))
    box(slide, x, y, w, 0.28, fill=fill)
    text(slide, x, y + 0.01, w, 0.28, label, size=9.5, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER)


def page(title: str, n: int):
    s = prs.slides.add_slide(BLANK)
    text(s, 0.35, 0.16, 12.6, 0.5, title, size=22, bold=True, color=NAVY, font="Cambria")
    box(s, 0, 7.06, 13.333, 0.44, fill=NAVY)
    text(s, 0.35, 7.13, 11.4, 0.35, FOOTER, size=9, color=GREY)
    text(s, 12.4, 7.11, 0.55, 0.35, str(n), size=11, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER)
    return s


def panel(s, x, y, w, h, label, body, *, fill=BLUE_BG, tagfill=BLUE, size=12.5):
    box(s, x, y, w, h, fill=fill)
    tag(s, x + 0.1, y - 0.13, label, fill=tagfill)
    text(s, x + 0.2, y + 0.25, w - 0.4, h - 0.35, body, size=size)


def stat(s, x, y, w, value, label, note, color=NAVY):
    text(s, x, y, w, 0.7, value, size=30, bold=True, color=color, font="Cambria")
    text(s, x, y + 0.72, w, 0.35, label, size=11.5, bold=True, color=INK)
    text(s, x, y + 1.05, w, 0.9, note, size=10.5, color=MUTED)


def table(s, x, y, w, rows, col_w, *, size=11.5, header_fill=NAVY):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                           Inches(0.36 * len(rows))).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = str(val)
            for para in cell.text_frame.paragraphs:  # every line, not just the first
                para.font.size = Pt(size)
                para.font.name = "Calibri"
                para.font.bold = i == 0
                para.font.color.rgb = WHITE if i == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (
                WHITE if i % 2 else RGBColor(0xF3, 0xF5, 0xF9))
    return t


def image(s, name, x, y, w=None, h=None, caption=None):
    path = IMG / name
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)
    pic.line.color.rgb = RGBColor(0xD0, 0xD5, 0xDD)
    if caption:
        bottom = (pic.top + pic.height) / 914400
        text(s, x, bottom + 0.05, pic.width / 914400, 0.3, caption, size=10, color=MUTED,
             italic=True)
    return pic


# ── 1 title ─────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, fill=NAVY)
text(s, 0.9, 2.1, 11.5, 1.0, "ATLAS", size=60, bold=True, color=WHITE, font="Cambria")
text(s, 0.9, 3.2, 11.5, 0.6, "Agentic Technology for Literature Analysis and Synthesis",
     size=24, color=GREY, font="Cambria")
text(s, 0.9, 4.1, 11.5, 0.9,
     "A literature-review agent that reads, follows citations, and keeps every claim\n"
     "tied to a verbatim, checked quote from its source.", size=16, color=WHITE)
text(s, 0.9, 5.55, 11.5, 0.45, AUTHORS, size=18, bold=True, color=WHITE)
text(s, 0.9, 6.3, 11.5, 0.4, "Final Review   ·   3 October 2026", size=13, color=GREY)


# ── 2 problem + the loop ────────────────────────────────────────────────────────
s = page("The Problem, and ATLAS's Answer: A Closed Foraging Loop", 2)
panel(s, 0.4, 0.95, 6.1, 1.65, "HOW TOOLS DO IT TODAY",
      "Search → take the top-k most similar abstracts → summarise. One pass: nothing read "
      "changes what is read next, and quotes can be paraphrased or invented.",
      fill=RED_BG, tagfill=RED, size=14)
panel(s, 6.8, 0.95, 6.1, 1.65, "HOW A RESEARCHER DOES IT",
      "Read the most promising paper → follow its references and citers → choose what to "
      "read next from what was just learned. Keep notes that point to the exact sentence.",
      fill=GREEN_BG, tagfill=GREEN, size=14)
steps = [("Harvest", "4 APIs → ~500\ncandidates"), ("Rank", "embed + score\nthe frontier"),
         ("Read", "PDF → GROBID\nsections"), ("Extract", "claims with\nverbatim quotes"),
         ("Ground", "quote must match\nsource exactly"), ("Judge", "claim-pair\nrelations"),
         ("Expand", "references +\ncitations"), ("Re-score", "what to read\nnext")]
for i, (h, sub) in enumerate(steps):
    x = 0.45 + i * 1.6
    box(s, x, 2.85, 1.42, 1.15, fill=BLUE_BG if i != 4 else AMBER_BG,
        line=BLUE if i != 4 else AMBER)
    text(s, x, 2.9, 1.42, 0.4, h, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x, 3.3, 1.42, 0.7, sub, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        text(s, x + 1.4, 3.15, 0.25, 0.4, "›", size=20, bold=True, color=NAVY)
text(s, 0.45, 4.1, 12.4, 0.4, "↺  repeat until the budget or saturation stops it  →  "
     "then: concepts → 5 gap detectors → one-pass grounded report → dashboard + chat",
     size=12.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
panel(s, 0.45, 4.8, 6.1, 1.95, "SCORING THE FRONTIER (what to read next)",
      "relevance · recency · novelty vs what was read · PageRank in the growing "
      "citation graph · contest (disputed claims)\n"
      "Bands on arrival: drop / park / queue; a parked paper bridging two clusters "
      "can be promoted.", size=13.5)
panel(s, 6.85, 4.8, 6.0, 1.95, "BUILT AND RUNNING",
      "Python · FastAPI · React + Cytoscape dashboard (live graph, report, chat)\n"
      "Neo4j · Qdrant · SQLite / Postgres · GROBID · local bge embeddings\n"
      "Free-tier LLMs: Groq gpt-oss (claims, judge), Gemini (report)",
      fill=GREEN_BG, tagfill=GREEN, size=13.5)

# ── 3 trust by construction ─────────────────────────────────────────────────────
s = page("Trust by Construction: Four Rules the Code Enforces", 3)
rules = [
    ("NO UNVERIFIED QUOTE", RED, RED_BG,
     "Every claim carries a quote that must match the normalised source text character for "
     "character. Failures are dropped and logged — never repaired by the LLM."),
    ("NO GAP WITHOUT CITATIONS", AMBER, AMBER_BG,
     "A research gap that cannot name the papers bounding it is filtered out before it is "
     "ranked, reported or shown."),
    ("NOTHING IT DID NOT READ", BLUE, BLUE_BG,
     "The chat retrieves only from this run's claims. With nothing relevant it refuses — "
     "decided by retrieval, before any LLM is called."),
    ("REPLAYABLE", GREEN, GREEN_BG,
     "Every LLM and API call is cached on its exact input. A re-run of the same query and "
     "config reproduces the run — measured, not assumed."),
]
for i, (label, c, bg, body) in enumerate(rules):
    x = 0.4 + (i % 2) * 6.35
    y = 0.95 + (i // 2) * 2.95
    panel(s, x, y, 6.1, 2.6, label, body, fill=bg, tagfill=c, size=14)

# ── 4 demo ──────────────────────────────────────────────────────────────────────
s = page("Demo — The Live Graph, and Every Claim Pointing to Its Sentence", 4)
image(s, "graph.png", 0.4, 0.9, w=6.1,
      caption="Run a0e32253 (IoT devices for women's safety), built live paper by paper.")
image(s, "paper_highlight.png", 6.8, 0.9, w=6.1,
      caption="Click a claim: its quote lights up inside the text ATLAS actually read.")
panel(s, 0.4, 4.85, 6.1, 1.85, "READING THE GRAPH",
      "Ringed nodes: the 15 papers read · faded: papers they cite.  Size = PageRank, "
      "colour = citation cluster.  Edges: grey cites · green supports · blue extends · "
      "dashed different scope · red contradicts.", size=13)
panel(s, 6.8, 4.85, 6.1, 1.85, "THE GROUNDING GATE, VISIBLE",
      "'Quote found verbatim in source' is recomputed in the browser from stored offsets.  "
      "Chat refuses 'Who won the 2018 FIFA World Cup?' — 'ATLAS did not read anything "
      "that addresses this.'", fill=GREEN_BG, tagfill=GREEN, size=13)

# ── 5 E1 ───────────────────────────────────────────────────────────────────────
s = page("Result 1 — Does the Loop Read Differently from One-Shot Retrieval?", 5)
table(s, 0.4, 0.95, 8.3, [
    ["Run (query, papers read)", "Overlap with\ncosine top-25", "Clusters\nATLAS / baseline*", "References reached\nATLAS / baseline*"],
    ["IoT safety, 15", "53%", "8 / 9", "115 / 92"],
    ["IoT safety, 15", "40%", "9 / 9", "147 / 92"],
    ["Federated learning, 14", "36%", "10 / 8", "355 / 332"],
    ["Federated learning, 13 (GROBID)", "38%", "11 / 9", "452 / 332"],
    ["Gut microbiome, 9 (GROBID)", "56%", "7 / 8", "381 / 365"],
], [3.0, 1.6, 1.8, 1.9], size=12)
text(s, 0.4, 3.62, 8.3, 0.4, "* baseline cut to the same number of papers; one OpenAlex citation "
     "graph and one Louvain clustering for both sets.", size=10.5, color=MUTED, italic=True)
panel(s, 9.0, 0.95, 3.9, 2.7, "HELD", "ATLAS reads a substantially different set (36–56% overlap) "
      "and reaches 4–60% more of the reference network.", fill=GREEN_BG, tagfill=GREEN, size=13)
panel(s, 9.0, 3.95, 3.9, 2.7, "NOT SHOWN", "At 9–15 papers it is not broader in citation-cluster "
      "coverage (−1, 0, +2, +2, −1). Bridge promotion never fired at this scale.",
      fill=AMBER_BG, tagfill=AMBER, size=13)
panel(s, 0.4, 4.25, 8.3, 2.4, "CORRECTION WE MADE",
      "An earlier draft claimed '14 clusters vs ~1 for the baseline'. That compared two different "
      "metrics (citation clusters vs text-similarity clusters), on vectors a bug had overwritten. "
      "We found it, fixed both, and withdrew the figure. The table is the like-for-like result.",
      fill=RED_BG, tagfill=RED, size=13)

# ── 6 grounding ─────────────────────────────────────────────────────────────────
s = page("Result 2 — Are the Quotes Real?", 6)
stat(s, 0.5, 1.0, 3.0, "49 / 50", "quotes verbatim (98%)",
     "Each source re-fetched and checked as an exact substring.")
stat(s, 3.7, 1.0, 3.0, "4", "live edits caught", "Quotes the old fuzzy match let through: "
     "'is set too high', 'a'→'an', a dropped 'that', a PDF page number.", color=RED)
stat(s, 6.9, 1.0, 3.0, "3 / 198", "claims dropped", "Federated-learning run with GROBID — "
     "the gate rejects little once text is clean.", color=GREEN)
stat(s, 10.1, 1.0, 2.9, "0", "repairs by the LLM", "A failed quote is dropped and logged, "
     "never rewritten.", color=NAVY)
panel(s, 0.4, 3.35, 6.1, 3.3, "THE ONE FAILURE, TRACED",
      "The quote said 'the last few 5 layers'. The paper says 'the last few layers'.\n\n"
      "Checking the PDF: page 4 ends mid-sentence, page 5 opens with its number '5'. "
      "The PDF parser kept it; the LLM quoted what ATLAS read, faithfully.\n\n"
      "Fixed at the source: page numbers are stripped.", size=12.5)
panel(s, 6.8, 3.35, 6.1, 3.3, "THE LOOPHOLE WE CLOSED",
      "A 97%-similarity fallback existed for PDF noise — but it also accepted "
      "'increased'→'decreased' or 'p > 0.05'→'p < 0.05'.\n\n"
      "Now a fuzzy match must be exact once only spacing, hyphenation, quote marks, "
      "case and reference markers are ignored. Checked against every stored fuzzy match "
      "before shipping.", fill=GREEN_BG, tagfill=GREEN, size=12.5)

# ── 7 relations + gaps ──────────────────────────────────────────────────────────
s = page("Result 3 — Relations and Gaps: The Weakest Parts, Measured", 7)
table(s, 0.4, 0.95, 6.1, [
    ["Compared (12 single-label pairs)", "Agreement", "Cohen's κ"],
    ["Human vs AI annotator", "2 / 12", "−0.17"],
    ["Human vs ATLAS judge", "2 / 12", "−0.05"],
    ["AI annotator vs ATLAS judge", "5 / 12", "0.22"],
], [3.3, 1.4, 1.4], size=12)
panel(s, 0.4, 2.75, 6.1, 3.95, "RELATIONS",
      "Even the annotators disagree on 'supports' vs 'extends' vs 'different scope' (the human "
      "marked 3 of 15 pairs undecided), so a single judge-accuracy figure (57% or 36%) is not "
      "trustworthy.\n\n"
      "IoT's only contradiction ('NodeMCU' vs 'Raspberry Pi Pico') was false. Gut microbiome's "
      "('repeatedly associated with depression' vs 'clinical results are mixed') is a plausible "
      "tension for a human to judge.\n\n"
      "A stricter judge prompt removed the false contradiction but lost 8 of 9 genuine "
      "'supports' — a trade-off, not a fix.", fill=RED_BG, tagfill=RED, size=13)
panel(s, 6.8, 0.95, 6.1, 1.75, "GAP QUALITY",
      "IoT (10 gaps): 2 plausible · 4 trivial · 4 wrong\n"
      "IoT (6 gaps): 0 plausible · 1 trivial · 5 wrong\n"
      "Most 'unlinked concepts' are linked under other wording. AI-annotator ratings.",
      fill=AMBER_BG, tagfill=AMBER, size=12.5)
table(s, 6.8, 2.95, 6.1, [
    ["Federated learning", "PyMuPDF", "GROBID"],
    ["Full-text ratio", "14%", "23%"],
    ["Claims extracted", "98", "198"],
    ["Gaps (5 detectors)", "0", "0"],
], [2.5, 1.8, 1.8], size=12)
text(s, 6.8, 4.55, 6.1, 2.1,
     "GROBID doubles the claims per full-text paper, but open PDFs are scarce (IoT: 0 of 15), "
     "so the full-text gap detectors are starved, not broken — the remedy is scale. "
     "Gut microbiome (GROBID): 3 of 9 in full text, 30 / 65 / 13 claims vs 0–6 per abstract; "
     "3 gaps survived (not yet rated).", size=12.5, color=INK)

# ── 8 reproducibility + cost ────────────────────────────────────────────────────
s = page("Result 4 — Reproducible and Cheap", 8)
stat(s, 0.5, 1.0, 3.0, "Identical", "replay of a 15-paper run",
     "Same papers in the same order, 65 claims, 132 relations — compared field by field.",
     color=GREEN)
stat(s, 3.7, 1.0, 3.0, "0", "new LLM calls on replay", "145 of 145 calls served from cache; 6.7 min.",
     color=GREEN)
stat(s, 6.9, 1.0, 3.0, "≤ $0.05", "per 15-paper run", "At paid Groq rates, from measured tokens; "
     "free on free tiers.", color=NAVY)
stat(s, 10.1, 1.0, 2.9, "1–4 h", "fresh run, free tier", "Limit: 8K tokens/min and 200K/day per "
     "model — throughput, not money.", color=AMBER)
panel(s, 0.4, 3.35, 12.5, 3.3, "REPRODUCIBILITY BUGS FOUND AND FIXED ALONG THE WAY",
      "Gap lists differed between two runs of the detectors on the same graph (6 vs 7 gaps) — Python's "
      "per-process set order steered clustering and tie-breaks; now byte-identical across processes.\n"
      "Reading a paper silently overwrote its shared embedding, changing the next run's reading order.\n"
      "Answers cut off by the output limit were never cached, so replays asked the LLM again.\n"
      "Per-run facts (read order, gaps, contradictions) leaked between runs through shared paper nodes.",
      fill=GREEN_BG, tagfill=GREEN, size=13)

# ── 9 limitations + conclusion ──────────────────────────────────────────────────
s = page("Limitations, Conclusion and Future Work", 9)
lim = [
    "Scale: 9–15 papers per run, three questions; a real review covers hundreds.",
    "Breadth: reads differently, but broader citation coverage not yet shown.",
    "Relations: annotators disagree on the labels; the IoT contradiction was false.",
    "Gaps: mostly wording artefacts; 0–2 of 6–10 rated plausible.",
    "Full text: open PDFs are scarce (0–33%); much is read from abstracts.",
    "Annotation: one human on 15 pairs; other ratings by an AI annotator.",
    "Not run in time on the free tier: ablation for queries 2–3, a 75-paper deep run.",
]
box(s, 0.4, 0.95, 6.1, 5.75, fill=RED_BG)
tag(s, 0.5, 0.82, "LIMITATIONS — STATED PLAINLY", fill=RED)
text(s, 0.6, 1.2, 5.7, 5.4, [[("•  ", {"bold": True, "color": RED}), (t, {})] for t in lim],
     size=15)
panel(s, 6.8, 0.95, 6.1, 2.75, "WHAT ATLAS IS",
      "A working end-to-end research prototype: it harvests, reads, follows citations, builds a "
      "live graph and a grounded report, and answers only from what it read.\n"
      "Dependable as an evidence finder — every claim traceable to a machine-checked sentence, "
      "runs reproducible. Its paper-to-paper relations are leads, not conclusions.",
      fill=GREEN_BG, tagfill=GREEN, size=14)
panel(s, 6.8, 4.0, 6.1, 2.7, "NEXT",
      "Run at scale (75–200 papers) to test breadth and bridge promotion.\n"
      "Sharpen the relation taxonomy; two or more human annotators; re-measure the judge.\n"
      "Normalise concept wording before gap detection.\n"
      "Paid or local LLM capacity to lift the free-tier ceiling.", size=14)

# ── 10 thank you ────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, fill=NAVY)
text(s, 0.9, 2.6, 11.5, 1.0, "Thank you", size=54, bold=True, color=WHITE, font="Cambria")
text(s, 0.9, 3.7, 11.5, 0.6, "Questions  ·  live demo on request", size=22, color=GREY,
     font="Cambria")
text(s, 0.9, 6.3, 11.5, 0.4, AUTHORS, size=13, color=GREY)

prs.save(OUT)
print(f"saved {OUT} ({len(prs.slides)} slides)")
